"""Readable MCP evaluation records. Completion never implies business accuracy."""
from __future__ import annotations

import html
import json
import os
from pathlib import Path
import shutil
import subprocess
import xml.etree.ElementTree as ET
import zipfile


def cell_text(value):
    text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    # Excel's limit is UTF-16 code units. Complete responses remain in HTML/JSON.
    if len(text.encode("utf-16-le")) // 2 > 32000:
        text = text.encode("utf-16-le")[:63000].decode("utf-16-le", errors="ignore") + "\n[超出Excel单元格容量，全文见完整阅读页/原始JSON]"
    return text


def spec(name, title, note, headers, widths, rows, formats=None, **kwargs):
    return {"name": name, "title": title, "context": "来源：本批MCP实际响应；原始记录见results.json。",
            "note": note, "headers": headers, "widths": widths, "formats": formats or {},
            "rows": rows, **kwargs}


def report_sheets(records, manifest, reviews=None):
    reviews = reviews or {}
    main, metrics, steps, calls, data = [], [], [], [], []
    for record in records:
        case = record["case"]
        cid = case["case_id"]
        payload = record.get("response") or {}
        diag = payload.get("detail") or {}
        execution = payload.get("execution") or {}
        usage = diag.get("usage") or {}
        request = record.get("request") or {}
        review = reviews.get(cid) or {}
        if isinstance(review, str):
            review = {"verdict": review}
        error = record.get("error") or record.get("mcp_error") or payload.get("error") or record.get("problems")
        seconds = record.get("elapsed_seconds")
        if seconds is None and record.get("client_elapsed_ms") is not None:
            seconds = record["client_elapsed_ms"] / 1000
        selected = request.get("skill_ids") or []
        tool = record.get("tool") or manifest.get("tool") or "finance_task"
        used = ", ".join(s.get("skill_id", "") for s in diag.get("skills") or [])
        main.append([cid, case["question"], cell_text(payload.get("summary") or ""), "打开全文",
            "指定：" + ", ".join(selected) if selected else ("自动选择" if tool == "finance_task" else "金融数据查询"),
            used or ("未返回" if not diag else "未记录Skill"), seconds, diag.get("turns"),
            diag.get("total_tokens"), "完成" if payload.get("ok") is True and not error else "未完成",
            review.get("verdict", "未人工评审"), review.get("evidence", ""), cell_text(error) if error else ""])
        metrics.append([cid, tool, execution.get("model_name"), execution.get("reasoning_effort"),
            diag.get("request_duration_ms"), execution.get("duration_ms"), execution.get("queue_wait_ms"),
            usage.get("prompt_tokens"), usage.get("cache_read_tokens"), usage.get("cumulative_context_tokens"),
            usage.get("completion_tokens"), usage.get("reasoning_tokens"), usage.get("accounting_total_tokens"),
            execution.get("tool_call_count"), len([s for s in diag.get("steps") or [] if s.get("kind") == "tool"]) if diag else None,
            execution.get("result_count"), execution.get("total_rows"), execution.get("returned_rows"),
            execution.get("truncated"), payload.get("created_at"), cell_text(request), cell_text(execution), cell_text(diag.get("usage") or {})])
        for s in diag.get("steps") or []:
            if s.get("kind") == "llm":
                u = s.get("usage") or {}
                steps.append([cid, s.get("request_index"), s.get("turn"), s.get("step"), s.get("duration_ms"),
                    u.get("input_tokens"), u.get("cache_read_tokens"), u.get("context_tokens"),
                    u.get("output_tokens"), u.get("reasoning_tokens"), u.get("non_reasoning_output_tokens"),
                    u.get("total_tokens"), s.get("timing_basis"), cell_text(s)])
            else:
                calls.append([cid, "事件", s.get("tool"), s.get("turn"), s.get("step"), s.get("duration_ms"),
                    None, s.get("is_error"), cell_text(s.get("arguments") or {}), cell_text(s)])
        for c in diag.get("tool_calls") or []:
            calls.append([cid, "业务明细", c.get("tool"), None, c.get("flow_step"), c.get("duration_ms"),
                c.get("row_count"), cell_text(c.get("error") or c.get("validation_errors") or c.get("execution_error") or ""),
                c.get("submitted_request") or c.get("request") or c.get("goal"), cell_text(c)])
        for r in (payload.get("data") or {}).get("results") or []:
            data.append([cid, r.get("result_name"), r.get("goal"), r.get("row_count"), r.get("rows_returned"),
                r.get("truncated"), cell_text(r.get("schema") or {}), cell_text(r.get("rows") or []),
                cell_text(payload.get("data_sources") or [])])
    scope = "服务器隔离测试" if manifest.get("isolated") else "MCP客户端评测（服务器版本须按运行证据核实）"
    result = spec("评测结果", "MCP评测｜一题一行", "长回答完整存于单元格；受Excel行高上限限制，点击完整阅读查看全文。接口完成不等于业务正确。",
        ["编号", "完整问题", "完整回答", "完整阅读", "请求选择", "实际Skill", "客户端秒", "模型响应次数", "总Token", "接口结果", "人工结论", "复核证据", "异常"],
        [24, 48, 100, 16, 32, 32, 15, 16, 17, 14, 40, 60, 40], main, {6:"0.00",7:"#,##0",8:"#,##0"},
        links={"3": [f"完整阅读.html#case-{i}" for i in range(len(records))]}, max_row_height=240)
    result["context"] = f"{scope} · {manifest.get('created_at') or manifest.get('started_at') or ''} · detail={manifest.get('detail', '按原始请求')}"
    return [result,
        spec("运行指标", "模型、Token与服务耗时", "时间单位ms；不同计时可能重叠，不相加。缓存与推理Token遵循API口径；缺失不记零。",
            ["编号","MCP入口","模型","推理强度","请求耗时ms","执行耗时ms","排队ms","非缓存输入Token","缓存输入Token","累计上下文Token","输出Token","推理Token","记账总Token","API工具计数","工具事件数","结果集数","结果总行数","返回行数","截断","API创建时间原值","请求参数","execution原值","usage原值"],
            [24,26,32,16,18,18,16,22,22,24,18,18,18,18,18,16,18,18,12,30,65,70,65], metrics,
            {**{i:"#,##0" for i in range(5,18)},4:"0.000"}, max_row_height=54),
        spec("模型步骤", "逐次模型响应", "每行一次模型响应；含推理Token与执行证据。当前API未返回内部思考原文，也未返回独立的纯思考耗时。",
            ["编号","请求序号","内部turn","内部step","耗时ms","非缓存输入","缓存输入","上下文输入","输出Token","推理Token","非推理输出","总Token","计时边界","步骤原值"],
            [24,14,14,14,16,18,18,18,18,18,18,18,60,90], steps,{i:"#,##0" for i in range(1,12)}, max_row_height=72),
        spec("工具调用", "工具事件与业务调用", "一行一次事件或业务调用；同一次调用可能分别有事件与业务记录，不能把两类行数直接相加。尝试、错误及耗时原值全部保留。",
            ["编号","记录来源","工具","内部turn","step/流程步","耗时ms","行数","错误/失败","参数或查询","完整记录（含尝试）"],
            [24,16,34,14,16,18,16,42,95,105],calls,{3:"0",4:"0",5:"0.000",6:"#,##0"}, max_row_height=110),
        spec("参考数据", "API返回的数据集", "每行一个结果集；保留接口返回样本、schema及来源。总行数含重复查询或聚合，不等于独立研报篇数。",
            ["编号","结果名","用途","总行数","返回行数","截断","Schema","实际返回数据","来源说明"],
            [24,18,60,16,16,12,70,110,80],data,{3:"#,##0",4:"#,##0"})]


def export_reading(records, manifest, folder, reviews):
    """No scripts or external resources; escape all model/source HTML."""
    esc = lambda v: html.escape(str(v if v is not None else ""))
    parts = ['<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>MCP评测完整阅读</title>',
        '<style>body{max-width:1060px;margin:32px auto;padding:0 24px;font:16px/1.7 system-ui;color:#243247}pre{white-space:pre-wrap;overflow-wrap:anywhere;font:inherit}article{margin:40px 0;border-top:2px solid #314763}details{background:#f1f4f8;padding:12px;margin:10px 0}a{color:#245faa}</style>',
        '<h1>MCP评测完整阅读</h1><p>原始回答未改写。接口完成不代表业务正确。</p><nav>']
    for i, r in enumerate(records):
        parts.append(f'<a href="#case-{i}">{esc(r["case"]["case_id"])}</a>　')
    parts.append('</nav><details><summary>本次运行配置</summary><pre>'+esc(json.dumps(manifest,ensure_ascii=False,indent=2))+'</pre></details>')
    for i, r in enumerate(records):
        c=r['case']; p=r.get('response') or {}
        parts.append(f'<article id="case-{i}"><h2>{esc(c["case_id"])}</h2><h3>问题</h3><pre>{esc(c["question"])}</pre><h3>原始回答</h3><pre>{esc(p.get("summary") or "未返回回答")}</pre>')
        for title, value in [('人工复核',reviews.get(c['case_id'], '未人工评审')),('执行证据与返回数据',r)]:
            parts.append(f'<details><summary>{title}</summary><pre>{esc(json.dumps(value,ensure_ascii=False,indent=2))}</pre></details>')
        parts.append('</article>')
    parts.append('</html>')
    (folder/'完整阅读.html').write_text(''.join(parts),encoding='utf-8')


def _artifact_runtime():
    node = os.environ.get("FINANCE_EVAL_NODE") or shutil.which("node")
    modules = os.environ.get("FINANCE_EVAL_ARTIFACT_MODULES")
    if not node:
        return None
    if modules:
        if not (Path(modules) / "@oai/artifact-tool").is_dir():
            raise RuntimeError("FINANCE_EVAL_ARTIFACT_MODULES does not contain @oai/artifact-tool")
        return node, Path(modules)
    check = subprocess.run([node, "--input-type=module", "-e",
        "try { console.log(import.meta.resolve('@oai/artifact-tool')) } catch { process.exit(3) }"],capture_output=True,text=True)
    return (node,None) if check.returncode == 0 else None


def export_report(records, manifest, folder, *, reviews=None):
    folder = Path(folder).resolve()
    sheets = report_sheets(records, manifest, reviews)
    export_reading(records, manifest, folder, reviews or {})
    grid = folder / "report_grid.json"
    grid.write_text(json.dumps(sheets, ensure_ascii=False, indent=2))
    output = folder / "MCP评测结果.xlsx"
    runtime = _artifact_runtime()
    if runtime:
        node, modules = runtime
        if modules:
            link = folder / "node_modules"
            if not link.exists():
                link.symlink_to(modules, target_is_directory=True)
        builder = folder / "export_excel.mjs"
        shutil.copyfile(Path(__file__).with_name("finance_mcp_report.mjs"), builder)
        completed = subprocess.run([node, str(builder), str(grid), str(output)],capture_output=True,text=True)
        if completed.returncode:
            raise RuntimeError("Artifact Tool export failed; JSON retained. Check exporter dependencies.")
        _add_reading_links(output, sheets)
    else:
        _export_python(sheets, output)
    return output


def _add_reading_links(output, sheets):
    """Native hyperlinks: Artifact Tool has no supported hyperlink setter/evaluator.

    Patch only OOXML link relationships; all cells, tables and styles are authored
    by Artifact Tool. This avoids unsupported HYPERLINK formula cached results.
    """
    sheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    doc_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(output) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    for i, s in enumerate(sheets, 1):
        if not s.get("links"):
            continue
        name = f"xl/worksheets/sheet{i}.xml"
        rel_name = f"xl/worksheets/_rels/sheet{i}.xml.rels"
        root = ET.fromstring(files[name])
        rels = ET.fromstring(files[rel_name]) if rel_name in files else ET.Element(f"{{{rel_ns}}}Relationships")
        links = ET.Element(f"{{{sheet_ns}}}hyperlinks")
        for column, targets in s["links"].items():
            for row, target in enumerate(targets, 7):
                rid = f"readingLink{column}_{row}"
                ET.SubElement(rels, f"{{{rel_ns}}}Relationship", {"Id":rid, "Type":doc_ns+"/hyperlink", "Target":target, "TargetMode":"External"})
                ET.SubElement(links, f"{{{sheet_ns}}}hyperlink", {"ref":f"{chr(65+int(column))}{row}", f"{{{doc_ns}}}id":rid})
        boundary = next((j for j, child in enumerate(root) if child.tag.rsplit('}',1)[-1] in
                        {"printOptions","pageMargins","pageSetup","headerFooter","drawing","tableParts","extLst"}),len(root))
        root.insert(boundary, links)
        files[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        files[rel_name] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
    temporary = output.with_suffix(".links.tmp")
    with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    temporary.replace(output)


def _export_python(sheets, output):
    """Standalone server fallback when Artifact Tool is not installed."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    wb = Workbook()
    wb.remove(wb.active)
    for si, s in enumerate(sheets):
        ws = wb.create_sheet(s['name']); ws.sheet_view.showGridLines=False
        for row, value in [(2,s['title']),(3,s['context']),(4,s['note'])]:
            ws.cell(row,1,value).font=Font(name='Arial',size=16 if row==2 else 11,bold=row==2)
        ws.row_dimensions[2].height=28
        for ri, values in enumerate([s['headers'],*s['rows']],6):
            for ci,v in enumerate(values,1):
                c=ws.cell(ri,ci,v)
                if isinstance(v,str):c.data_type='s'
                c.font=Font(name='Arial',size=11,color='FFFFFF' if ri==6 else '253449',bold=ri==6)
                c.alignment=Alignment(vertical='top',wrap_text=True)
                if ri==6:c.fill=PatternFill('solid',fgColor='24476B')
                if ri>6 and ci-1 in s['formats']:c.number_format=s['formats'][ci-1]
                if ri>6 and str(ci-1) in s.get('links',{}):
                    c.hyperlink=s['links'][str(ci-1)][ri-7];c.font=Font(name='Arial',size=11,color='245FAA',underline='single')
            ws.row_dimensions[ri].height=32 if ri==6 else min(s.get('max_row_height',300),max(36,max(
                sum(max(1,(sum(2 if ord(c)>255 else 1 for c in line)+w-4)//(w-3)) for line in str(v if v is not None else '').split('\n'))
                for v,w in zip(values,s['widths']))*15+12))
        for ci,w in enumerate(s['widths'],1):ws.column_dimensions[get_column_letter(ci)].width=w
        ws.freeze_panes='B7'
        if s['rows']:
            t=Table(displayName=f'McpResults{si}',ref=f'A6:{get_column_letter(len(s["headers"]))}{len(s["rows"])+6}')
            t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(t)
    wb.save(output)
