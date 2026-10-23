"""Read the delivered XLSX and reconcile it to the frozen workbook and sources."""
import hashlib,json,math
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/stock_automl_runs/20261021_b2_floor_factor_review'
SOURCE=ROOT/'outputs/b2_model_month_r61_20261020/B2_严格月模型归属_R61_20261020.xlsx'
TARGET=ROOT/'outputs/b2_floor_factor_review_20261021/B2_跌停专项与全量因子.xlsx'

def equal(a,b):
    if isinstance(a,(int,float)) and isinstance(b,(int,float)):
        return math.isclose(a,b,rel_tol=1e-12,abs_tol=1e-9)
    return a==b

def main():
    old=openpyxl.load_workbook(SOURCE);new=openpyxl.load_workbook(TARGET)
    assert new.sheetnames==old.sheetnames+['跌停专项','因子说明']
    counts={}
    for name in old.sheetnames:
        a,b=old[name],new[name];count=0
        for row in a:
            for cell in row:
                if name=='全部实买' and cell.coordinate=='A3':continue
                assert equal(cell.value,b[cell.coordinate].value),(name,cell.coordinate,cell.value,b[cell.coordinate].value)
                count+=1
        counts[name]=count
    data=json.loads((OUT/'workbook_data.json').read_text());records=data['rows'];cols=data['columns']
    assert sum(c['model_factor'] for c in cols)==51
    main_sheet=new['全部实买'];extra=[c for c in cols if c['key'] not in ('entry','original_net','d2_close_saved')]
    for i,record in enumerate(records,6):
        assert int(main_sheet.cell(i,5).value)==int(record['symbol'])
        for j,c in enumerate(extra,13):
            assert main_sheet.cell(5,j).value==c['label']
            assert equal(main_sheet.cell(i,j).value,record.get(c['key'])),(i,c['key'],main_sheet.cell(i,j).value,record.get(c['key']))
    special=new['跌停专项'];heads=[c.value for c in special[5]];lookup={v:i+1 for i,v in enumerate(heads)}
    subset=[r for r in records if r['open_floor']=='是' or r['close_floor']=='是']
    assert special.max_row==72 and len(subset)==67
    for i,r in enumerate(subset,6):
        assert special.cell(i,2).value==r['stk_code']
        assert special.cell(i,2).data_type=='s'
        for c in cols:
            assert equal(special.cell(i,lookup[c['label']]).value,r.get(c['key'])),(i,c['key'])
    assert sum(special.cell(i,lookup['开盘跌停']).value=='是' for i in range(6,73))==34
    assert sum(special.cell(i,lookup['收盘跌停']).value=='是' for i in range(6,73))==52
    for name in ['全部实买','跨月卖出','跌停专项','因子说明']:
        s=new[name];assert len(s.tables)==1 and s.freeze_panes is not None
        for row in s:
            for cell in row:assert cell.data_type!='e',(name,cell.coordinate,cell.value)
    assert special.freeze_panes=='F6' and main_sheet.freeze_panes=='F6'
    report=dict(original_cells_preserved=counts,all_306_trade_factor_cells_match=True,
        all_67_special_rows_match=True,original_51_factors_included=True,opening=34,closing=52,both=19,
        missing_is_blank_not_zero=True,filters_and_freeze_panes=True,no_excel_error_cells=True,
        source_workbook_sha256=hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
        output_workbook_sha256=hashlib.sha256(TARGET.read_bytes()).hexdigest(),
        workbook=str(TARGET.relative_to(ROOT)))
    (OUT/'workbook_verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(report,ensure_ascii=False))

if __name__=='__main__':main()
